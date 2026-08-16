"""データ構造の正規化と Excel 入出力を扱う。"""
from __future__ import annotations
import copy
import json
import re
import uuid
from pathlib import Path
from typing import Any
from core.data_values import normalize_data_record, strip_unicode_whitespace
from core.data_templates import (
    TEMPLATE_INSTANCES_KEY, iter_template_instances, normalize_template_schema,
    schema_templates, template_by_id,
)
from i18n import SUPPORTED_LANGUAGES, tr, tr_language
TYPES = ('text', 'number', 'boolean', 'object', 'list')

def strip_data_whitespace(value: Any) -> Any:
    """Data 管理で扱う文字列から先頭・末尾の Unicode 空白を除去する。"""
    return strip_unicode_whitespace(value)

def strip_data_record_whitespace(record: dict[str, Any]) -> dict[str, Any]:
    """Data の入出力用レコードに同じ空白除去規則を適用する。"""
    return normalize_data_record(record)

def strip_schema_name_whitespace(schema: dict[str, Any]) -> dict[str, Any]:
    """構造内の全フィールド名から先頭・末尾の Unicode 空白を除去する。"""
    cleaned = copy.deepcopy(schema)
    if not isinstance(cleaned, dict):
        validate_schema(cleaned)

    def walk(node: dict[str, Any]) -> None:
        node['name'] = str(node.get('name', '')).strip()
        for child in node.get('children', []):
            if isinstance(child, dict):
                walk(child)
    walk(cleaned)
    for template in schema_templates(cleaned):
        walk(template)
    validate_schema(cleaned)
    for template in schema_templates(cleaned):
        validate_schema(template, f"Template.{template.get('name', '')}")
    return cleaned

def remap_data_for_schema_names(
        old_schema: dict[str, Any], new_schema: dict[str, Any], data: Any) -> Any:
    """空白除去で変更された構造名に合わせ、既存データのキーを移行する。"""
    if old_schema.get('type') == 'list':
        if not isinstance(data, list):
            return data
        return [remap_data_for_schema_names(
            {'type': 'object', 'children': old_schema.get('children', [])},
            {'type': 'object', 'children': new_schema.get('children', [])},
            item,
        ) for item in data]
    if old_schema.get('type') != 'object' or not isinstance(data, dict):
        return data
    result = dict(data)
    for old_child, new_child in zip(old_schema.get('children', []), new_schema.get('children', [])):
        old_name, new_name = old_child['name'], new_child['name']
        if old_name not in result:
            continue
        value = result.pop(old_name)
        result[new_name] = remap_data_for_schema_names(old_child, new_child, value)
    return result

def schema_name_path_map(
        old_schema: dict[str, Any], new_schema: dict[str, Any]) -> dict[str, str]:
    """空白除去前後で変更された全構造パスの対応表を作成する。"""
    result: dict[str, str] = {}

    def walk(old_node: dict[str, Any], new_node: dict[str, Any], old_prefix: str, new_prefix: str) -> None:
        for old_child, new_child in zip(old_node.get('children', []), new_node.get('children', [])):
            old_path = f"{old_prefix}.{old_child['name']}" if old_prefix else old_child['name']
            new_path = f"{new_prefix}.{new_child['name']}" if new_prefix else new_child['name']
            if old_path != new_path:
                result[old_path] = new_path
            walk(old_child, new_child, old_path, new_path)
    walk(old_schema, new_schema, '', '')
    return result

def validate_schema(node: Any, location: str='Data', *, split_ancestor: bool=False) -> None:
    if not isinstance(node, dict) or not isinstance(node.get('name'), str) or (not node['name'].strip()):
        raise ValueError(f'{location}{tr("schema.invalid_name_suffix")}')
    if node.get('type') not in TYPES:
        raise ValueError(f'{location}{tr("schema.invalid_type_prefix")}{node.get("type")}')
    split_here = bool(node.get('excel_sheet', False))
    if split_here and node.get('type') != 'object':
        raise ValueError(f'{location}{tr(": Excel の別シート出力は object だけに設定できます。")}')
    if split_here and split_ancestor:
        raise ValueError(f'{location}{tr(": 別シート出力の配下では、別のシートを設定できません。")}')
    children = node.get('children', [])
    if node['type'] in ('object', 'list'):
        if not isinstance(children, list):
            raise ValueError(f'{location}{tr("schema.children_array_suffix")}')
        names: set[str] = set()
        for child in children:
            child_name = child.get('name') if isinstance(child, dict) else '?'
            if child_name in names:
                raise ValueError(f'{location}{tr("schema.duplicate_field_prefix")}{child_name}')
            names.add(child_name)
            validate_schema(
                child, f'{location}.{child_name}',
                split_ancestor=split_ancestor or split_here,
            )
    elif 'children' in node and children:
        raise ValueError(f'{location}{tr("schema.scalar_children_suffix")}')

def scalar_paths(schema: dict[str, Any]) -> list[str]:
    return [path for path, node in schema_paths(schema) if node['type'] not in ('object', 'list')]

def _flatten_record_with_groups(schema: dict[str, Any], data: dict[str, Any]) -> list[tuple[dict[str, Any], dict[str, int]]]:
    # list の添字も返し、同じ親値を Excel 上で何度も出力しないようにする。

    def merge(left: list[tuple[dict[str, Any], dict[str, int]]], right: list[tuple[dict[str, Any], dict[str, int]]]) -> list[tuple[dict[str, Any], dict[str, int]]]:
        return [({**a_values, **b_values}, {**a_groups, **b_groups}) for a_values, a_groups in left for b_values, b_groups in right]

    def walk_fields(nodes: list[dict[str, Any]], value: dict[str, Any], prefix: str) -> list[tuple[dict[str, Any], dict[str, int]]]:
        rows: list[tuple[dict[str, Any], dict[str, int]]] = [({}, {})]
        for node in nodes:
            path = f"{prefix}.{node['name']}" if prefix else node['name']
            current = value.get(node['name'])
            if node['type'] == 'list':
                items = current if isinstance(current, list) else []
                child_rows: list[tuple[dict[str, Any], dict[str, int]]] = []
                for item_index, item in enumerate(items):
                    for child_values, child_groups in walk_fields(node.get('children', []), item if isinstance(item, dict) else {}, path):
                        child_rows.append((child_values, {path: item_index, **child_groups}))
                if not child_rows:
                    child_rows = [({child_path: '' for child_path in _descendant_scalar_paths(node, path)}, {path: -1})]
            elif node['type'] == 'object':
                child_rows = walk_fields(node.get('children', []), current if isinstance(current, dict) else {}, path)
            else:
                child_rows = [({path: current if current is not None else ''}, {})]
            rows = merge(rows, child_rows)
        return rows
    return walk_fields(schema.get('children', []), data, '')

def scalar_list_owners(schema: dict[str, Any]) -> dict[str, tuple[str, ...]]:
    result: dict[str, tuple[str, ...]] = {}

    def walk(node: dict[str, Any], prefix: str, list_ancestors: tuple[str, ...]) -> None:
        for child in node.get('children', []):
            path = f"{prefix}.{child['name']}" if prefix else child['name']
            ancestors = (*list_ancestors, path) if child['type'] == 'list' else list_ancestors
            if child['type'] in ('object', 'list'):
                walk(child, path, ancestors)
            else:
                result[path] = list_ancestors
    walk(schema, '', ())
    return result

def _descendant_scalar_paths(node: dict[str, Any], prefix: str) -> list[str]:
    result: list[str] = []
    for child in node.get('children', []):
        path = f"{prefix}.{child['name']}"
        if child['type'] in ('object', 'list'):
            result.extend(_descendant_scalar_paths(child, path))
        else:
            result.append(path)
    return result

def _data_excel_sheet(workbook: Any) -> Any:
    """表示言語に依存せず、データ本体のシートを取得する。"""
    data_names = {
        tr_language('data_excel.data_sheet', language)
        for language in SUPPORTED_LANGUAGES
    }
    data_name = next((name for name in workbook.sheetnames if name in data_names), None)
    return workbook[data_name] if data_name is not None else workbook.worksheets[0]

def _excel_header_layout(sheet: Any) -> tuple[int, list[str]]:
    """結合された多段ヘッダーを、データ構造のパスへ展開する。"""
    header_depth = 1
    for merged in sheet.merged_cells.ranges:
        if merged.min_col == merged.max_col == 1 and merged.min_row == 1:
            header_depth = max(header_depth, merged.max_row)

    def header_value(row: int, column: int) -> Any:
        value = sheet.cell(row, column).value
        if value is not None:
            return value
        for merged in sheet.merged_cells.ranges:
            if merged.min_row <= row <= merged.max_row and merged.min_col <= column <= merged.max_col:
                # 縦結合は末端項目、横結合は配下列へ引き継ぐ親項目として扱う。
                if merged.min_row < row and merged.max_row > merged.min_row:
                    return None
                return sheet.cell(merged.min_row, merged.min_col).value
        return None

    columns: list[str] = []
    for column in range(2, sheet.max_column + 1):
        parts = [
            str(value).strip()
            for row in range(1, header_depth + 1)
            if (value := header_value(row, column)) not in (None, '')
        ]
        columns.append('.'.join(parts))
    return header_depth, columns

def _infer_schema_from_workbook(workbook: Any) -> dict[str, Any]:
    """Excel の多段ヘッダーと値から、データ構造を復元する。"""
    sheet = _data_excel_sheet(workbook)
    header_depth, columns = _excel_header_layout(sheet)
    if not columns or any(not path_name for path_name in columns):
        raise ValueError(tr('data_excel.header_mismatch'))
    if len(columns) != len(set(columns)):
        raise ValueError(tr('data_excel.duplicate_headers'))

    root: dict[str, Any] = {'name': 'Data', 'type': 'object', 'children': []}
    nodes_by_path: dict[str, dict[str, Any]] = {}
    for path_name in columns:
        parent = root
        prefix = ''
        parts = path_name.split('.')
        for index, name in enumerate(parts):
            prefix = f'{prefix}.{name}' if prefix else name
            node = nodes_by_path.get(prefix)
            if node is None:
                node = {
                    'name': name,
                    'type': 'text' if index == len(parts) - 1 else 'object',
                }
                if index < len(parts) - 1:
                    node['children'] = []
                parent.setdefault('children', []).append(node)
                nodes_by_path[prefix] = node
            elif index < len(parts) - 1 and node['type'] not in ('object', 'list'):
                raise ValueError(tr('data_excel.header_mismatch'))
            parent = node

    # Excel のセル型から基本型を決定する。混在する列は安全側で文字列として扱う。
    for column_index, path_name in enumerate(columns, 2):
        values = [
            sheet.cell(row, column_index).value
            for row in range(header_depth + 1, sheet.max_row + 1)
            if sheet.cell(row, column_index).value not in (None, '')
        ]
        node = nodes_by_path[path_name]
        if values and all(isinstance(value, bool) for value in values):
            node['type'] = 'boolean'
        elif values and all(
            isinstance(value, (int, float)) and not isinstance(value, bool)
            for value in values
        ):
            node['type'] = 'number'

    record_rows: list[list[int]] = []
    current_rows: list[int] = []
    for row in range(header_depth + 1, sheet.max_row + 1):
        if sheet.cell(row, 1).value not in (None, ''):
            if current_rows:
                record_rows.append(current_rows)
            current_rows = []
        if current_rows or sheet.cell(row, 1).value not in (None, ''):
            current_rows.append(row)
    if current_rows:
        record_rows.append(current_rows)

    # 同一データ内で直下項目が複数行に現れる階層はリストとして復元する。
    column_indexes = {path_name: index for index, path_name in enumerate(columns, 2)}
    for path_name, node in nodes_by_path.items():
        if node['type'] != 'object':
            continue
        direct_scalar_paths = [
            child_path for child_path, child in nodes_by_path.items()
            if child_path.rpartition('.')[0] == path_name
            and child['type'] not in ('object', 'list')
        ]
        if any(
            sum(
                sheet.cell(row, column_indexes[child_path]).value not in (None, '')
                for row in rows
            ) > 1
            for child_path in direct_scalar_paths
            for rows in record_rows
        ):
            node['type'] = 'list'
    validate_schema(root)
    return root

def read_records_excel_with_schema(
        path: str | Path,
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    """Excel からデータ構造を復元し、その構造でデータも読み込む。"""
    from openpyxl import load_workbook
    workbook = load_workbook(path, data_only=True)
    metadata = _excel_metadata(workbook)
    schema = normalize_template_schema(
        metadata['schema'] if metadata is not None else _infer_schema_from_workbook(workbook)
    )
    validate_schema(schema)
    return schema, read_records_excel(path, schema, workbook=workbook)

EXCEL_METADATA_SHEET = '__WebFlowManager__'
EXCEL_FORMAT_VERSION = 2


def _excel_metadata(workbook: Any) -> dict[str, Any] | None:
    """本アプリが出力した分割シート情報を、非表示シートから取得する。"""
    if EXCEL_METADATA_SHEET not in workbook.sheetnames:
        return None
    sheet = workbook[EXCEL_METADATA_SHEET]
    values = {
        str(sheet.cell(row, 1).value or ''): sheet.cell(row, 2).value
        for row in range(1, sheet.max_row + 1)
    }
    try:
        version = int(values.get('format_version', 0))
        schema = json.loads(str(values['schema_json']))
        mappings = json.loads(str(values['sheet_map_json']))
        template_presence = json.loads(str(values.get('template_presence_json', '{}')))
    except (KeyError, TypeError, ValueError, json.JSONDecodeError) as error:
        raise ValueError(tr('Excel の管理情報が正しくありません。')) from error
    if (
            version != EXCEL_FORMAT_VERSION
            or not isinstance(mappings, list)
            or not isinstance(template_presence, dict)
    ):
        raise ValueError(tr('対応していない Excel 形式です。'))
    return {'schema': schema, 'mappings': mappings, 'template_presence': template_presence}


def _split_excel_schema(
        schema: dict[str, Any],
) -> tuple[dict[str, Any], list[tuple[str, dict[str, Any]]]]:
    """別シート指定ノードを本体構造から除き、出力単位として収集する。"""
    sections: list[tuple[str, dict[str, Any]]] = []

    def clone(node: dict[str, Any], prefix: str) -> dict[str, Any]:
        copied = {key: copy.deepcopy(value) for key, value in node.items() if key != 'children'}
        children = []
        for child in node.get('children', []):
            child_path = f"{prefix}.{child['name']}" if prefix else child['name']
            if child.get('excel_sheet', False):
                sections.append((child_path, copy.deepcopy(child)))
            else:
                children.append(clone(child, child_path))
        if node.get('type') in ('object', 'list'):
            copied['children'] = children
        return copied

    return clone(schema, ''), sections


def _excel_sheet_name(name: str, used: set[str]) -> str:
    """Excel の禁止文字と31文字制限を処理し、重複しないシート名を返す。"""
    base = re.sub(r'[\\/*?:\[\]]', '_', name).strip(" '") or 'Data'
    base = base[:31]
    candidate = base
    suffix = 2
    while candidate.casefold() in used:
        tail = f'_{suffix}'
        candidate = f'{base[:31 - len(tail)]}{tail}'
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def _value_at_path(data: dict[str, Any], path_name: str) -> Any:
    value: Any = data
    for part in path_name.split('.'):
        if not isinstance(value, dict):
            return None
        value = value.get(part)
    return value


def _without_template_instances(value: Any) -> Any:
    """通常データ用シートからテンプレート実体だけを除外する。"""
    if isinstance(value, list):
        return [
            _without_template_instances(item) for item in value
            if not (
                isinstance(item, dict) and isinstance(item.get('data'), dict)
                and bool(str(item.get('template_id', '')).strip())
            )
        ]
    if isinstance(value, dict):
        return {
            key: _without_template_instances(child)
            for key, child in value.items() if key != TEMPLATE_INSTANCES_KEY
        }
    return copy.deepcopy(value)


def _assign_value_at_path(data: dict[str, Any], path: list[Any], value: Any) -> None:
    """文字列・配列添字が混在する保存パスへ値を復元する。"""
    current: Any = data
    for offset, part in enumerate(path):
        last = offset == len(path) - 1
        next_part = None if last else path[offset + 1]
        if isinstance(part, int):
            if not isinstance(current, list) or part < 0:
                raise ValueError('Invalid template container path')
            while len(current) <= part:
                current.append(None)
            if last:
                current[part] = value
            else:
                expected_type = list if isinstance(next_part, int) else dict
                if not isinstance(current[part], expected_type):
                    current[part] = expected_type()
                current = current[part]
        else:
            if not isinstance(current, dict):
                raise ValueError('Invalid template container path')
            key = str(part)
            if last:
                current[key] = value
            else:
                expected_type = list if isinstance(next_part, int) else dict
                if not isinstance(current.get(key), expected_type):
                    current[key] = expected_type()
                current = current[key]


def _display_container_path(path: tuple[Any, ...]) -> str:
    """配列位置を 1 始まりで示す、人が読める配置先を作成する。"""
    parts: list[str] = []
    for part in path:
        if isinstance(part, int) and parts:
            parts[-1] = f'{parts[-1]}[{part + 1}]'
        elif isinstance(part, str) and part not in {TEMPLATE_INSTANCES_KEY, 'data'}:
            parts.append(part)
    return '.'.join(parts)


def _resolve_container_path(data: dict[str, Any], display_path: str) -> list[Any]:
    """「親[1].子」形式の配置先から対象 list を取得する。"""
    current: Any = data
    for offset, token in enumerate(display_path.split('.')):
        match = re.fullmatch(r'(.+?)(?:\[(\d+)\])?', token)
        if match is None or not isinstance(current, dict):
            raise ValueError(f'配置先が正しくありません: {display_path}')
        name, index_text = match.groups()
        last = offset == len(display_path.split('.')) - 1
        current = current.setdefault(name, [] if last and index_text is None else {})
        if index_text is not None:
            if not isinstance(current, list):
                raise ValueError(f'配置先が list ではありません: {display_path}')
            index = int(index_text) - 1
            if index < 0 or index >= len(current):
                raise ValueError(f'配置先の順番が範囲外です: {display_path}')
            current = current[index]
    if not isinstance(current, list):
        raise ValueError(f'配置先が list ではありません: {display_path}')
    return current


def _write_records_sheet(sheet: Any, schema: dict[str, Any], records: list[dict[str, Any]]) -> int:
    """一つの構造範囲を、従来と同じ階層ヘッダー形式でシートへ出力する。"""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    columns = scalar_paths(schema)
    owners = scalar_list_owners(schema)
    max_depth = max((len(column.split('.')) for column in columns), default=1)
    sheet.cell(1, 1, tr('data_excel.record_name'))
    if max_depth > 1:
        sheet.merge_cells(start_row=1, start_column=1, end_row=max_depth, end_column=1)
    path_parts = [column.split('.') for column in columns]
    for column_index, parts in enumerate(path_parts, 2):
        for level, part in enumerate(parts, 1):
            sheet.cell(level, column_index, part)
        if len(parts) < max_depth:
            sheet.merge_cells(start_row=len(parts), start_column=column_index, end_row=max_depth, end_column=column_index)
    for level in range(1, max_depth + 1):
        start = 0
        while start < len(columns):
            prefix = tuple(path_parts[start][:level]) if len(path_parts[start]) >= level else None
            end = start + 1
            while end < len(columns) and prefix is not None and (tuple(path_parts[end][:level]) == prefix):
                end += 1
            if prefix is not None and end - start > 1:
                sheet.merge_cells(start_row=level, start_column=start + 2, end_row=level, end_column=end + 1)
            start = end
    output_row = max_depth + 1
    for record in records:
        clean_data = strip_data_whitespace(normalize_record(schema, record['data']))
        rows = _flatten_record_with_groups(schema, clean_data) or [({}, {})]
        seen_scopes: set[tuple[Any, ...]] = set()
        for row_index, (values, groups) in enumerate(rows):
            sheet.cell(output_row, 1, str(record['name']).strip() if row_index == 0 else '')
            for column_index, column in enumerate(columns, 2):
                ancestors = owners[column]
                scope = (column, *(groups.get(path, -1) for path in ancestors))
                if scope not in seen_scopes:
                    sheet.cell(output_row, column_index, values.get(column, ''))
                    seen_scopes.add(scope)
            output_row += 1
    header_level_colors = ('2F5597', '4472C4', '5B9BD5', '6F9FD5', '7EA6D8')
    for level, row in enumerate(sheet.iter_rows(min_row=1, max_row=max_depth)):
        color = header_level_colors[min(level, len(header_level_colors) - 1)]
        for cell in row:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor=color)
    sheet.freeze_panes = f'A{max_depth + 1}'
    headers = [tr('data_excel.record_name'), *columns]
    for index, header in enumerate(headers, 1):
        values = [str(sheet.cell(row, index).value or '') for row in range(1, sheet.max_row + 1)]
        sheet.column_dimensions[get_column_letter(index)].width = min(max(len(header) + 2, *(len(value) + 2 for value in values)), 45)
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                cell.alignment = Alignment(horizontal='left', vertical='center')
    return sheet.max_row - max_depth


def write_records_excel(path: str | Path, schema: dict[str, Any], records: list[dict[str, Any]]) -> int:
    """PCL と実行設定を、再読込可能な Excel ブックへ保存する。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    original_schema = normalize_template_schema(schema)
    schema = strip_schema_name_whitespace(original_schema)
    migrated_records = [
        record | {'data': remap_data_for_schema_names(original_schema, schema, record['data'])}
        for record in records
    ]
    main_schema, sections = _split_excel_schema(schema)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = tr('data_excel.data_sheet')
    main_records = [
        record | {'data': _without_template_instances(record.get('data', {}))}
        for record in migrated_records
    ]
    row_count = _write_records_sheet(sheet, main_schema, main_records)
    used_names = {name.casefold() for name in workbook.sheetnames}
    for path_name, node in sections:
        section_schema = {
            'name': 'Data', 'type': 'object',
            'children': copy.deepcopy(node.get('children', [])),
        }
        section_records = [
            record | {'data': (
                copy.deepcopy(value) if isinstance(value := _value_at_path(record['data'], path_name), dict)
                else {}
            )}
            for record in migrated_records
            if (
                not node.get('data_template', False)
                or isinstance(_value_at_path(record['data'], path_name), dict)
            )
        ]
        sheet_name = _excel_sheet_name(node['name'], used_names)
        section_default = default_value(node)
        has_data = any(
            strip_data_whitespace(normalize_record(section_schema, record['data'])) != section_default
            for record in section_records
        )
        if node.get('excel_skip_empty', True) and not has_data:
            continue
        section_sheet = workbook.create_sheet(sheet_name)
        _write_records_sheet(section_sheet, section_schema, section_records)
    # テンプレートは定義ごとに一枚のシートへまとめ、同じ PCL の複数実体を複数行で保持する。
    for template in schema_templates(schema):
        template_id = str(template.get('template_id', ''))
        template_schema = {
            'name': 'Data', 'type': 'object',
            'children': [
                {'name': '配置先', 'type': 'text'},
                {'name': '順番', 'type': 'number'},
                {'name': 'テンプレート名', 'type': 'text'},
                *copy.deepcopy(template.get('children', [])),
            ],
        }
        template_records: list[dict[str, Any]] = []
        for record in migrated_records:
            for instance_path, instance in iter_template_instances(record.get('data', {})):
                if str(instance.get('template_id', '')) != template_id:
                    continue
                values = copy.deepcopy(instance.get('data', {})) if isinstance(instance.get('data'), dict) else {}
                is_top_level = bool(instance_path and instance_path[0] == TEMPLATE_INSTANCES_KEY)
                position = instance_path[-1] + 1 if instance_path and isinstance(instance_path[-1], int) else 1
                values['配置先'] = '' if is_top_level else _display_container_path(instance_path[:-1])
                values['順番'] = position
                values['テンプレート名'] = str(instance.get('name', ''))
                template_records.append(record | {'data': values})
        sheet_name = _excel_sheet_name(str(template.get('name', '')), used_names)
        if template_records:
            template_sheet = workbook.create_sheet(sheet_name)
            _write_records_sheet(template_sheet, template_schema, template_records)
    settings_sheet = workbook.create_sheet(tr('data_excel.settings_sheet'))
    settings_sheet.append([tr('data_excel.record_name'), tr('common.summary'), tr('data_excel.enabled'), tr('data_excel.execution_group')])
    for record in migrated_records:
        settings_sheet.append([
            str(record['name']).strip(),
            str(record.get('summary', '')).strip(),
            bool(record.get('enabled', True)),
            str(record.get('execution_group', '1')).strip(),
        ])
    for current_sheet in workbook.worksheets:
        for row in current_sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cell.alignment = Alignment(horizontal='left', vertical='center')
    workbook.save(path)
    return row_count

def read_records_excel(
        path: str | Path, schema: dict[str, Any], *, workbook: Any | None=None,
        _sheet: Any | None=None, _read_settings: bool=True, _allow_split: bool=True,
) -> list[dict[str, Any]]:
    """write_records_excel が作成した階層ブックを PCL に復元する。"""
    from openpyxl import load_workbook
    schema = strip_schema_name_whitespace(schema)
    workbook = workbook or load_workbook(path, data_only=True)
    metadata = _excel_metadata(workbook) if _allow_split else None
    if metadata is None and _allow_split:
        # 現行形式は管理シートを持たず、現在の構造と表示シート名から対応を判定する。
        used_names = {
            name.casefold() for name in workbook.sheetnames
            if name in {tr_language('data_excel.data_sheet', language) for language in SUPPORTED_LANGUAGES}
        }
        mappings: list[dict[str, Any]] = []
        _main_schema, sections = _split_excel_schema(schema)
        for path_name, node in sections:
            sheet_name = _excel_sheet_name(str(node.get('name', '')), used_names)
            mappings.append({'sheet': sheet_name, 'path': path_name})
        for template in schema_templates(schema):
            sheet_name = _excel_sheet_name(str(template.get('name', '')), used_names)
            if sheet_name in workbook.sheetnames:
                mappings.append({
                    'sheet': sheet_name, 'template_id': str(template.get('template_id', '')),
                    'kind': 'template', 'simple': True,
                })
        if mappings:
            metadata = {'mappings': mappings, 'template_presence': {}}
    if metadata is not None:
        return _read_split_records_excel(path, workbook, schema, metadata)
    data_names = {tr_language('data_excel.data_sheet', language) for language in SUPPORTED_LANGUAGES}
    # Excel は最後に表示していたシートを active として保存するため、名前でデータシートを特定する。
    # シート名を持たない旧形式だけは、従来どおり先頭シートへフォールバックする。
    data_name = next((name for name in workbook.sheetnames if name in data_names), None)
    sheet = _sheet or (workbook[data_name] if data_name is not None else workbook.worksheets[0])
    header_depth = 1
    for merged in sheet.merged_cells.ranges:
        # 先頭列の縦結合範囲が、階層ヘッダーの深さを表す。
        if merged.min_col == merged.max_col == 1 and merged.min_row == 1:
            header_depth = max(header_depth, merged.max_row)

    def header_value(row: int, column: int) -> Any:
        value = sheet.cell(row, column).value
        if value is not None:
            return value
        for merged in sheet.merged_cells.ranges:
            if merged.min_row <= row <= merged.max_row and merged.min_col <= column <= merged.max_col:
                # 縦結合された末端項目は複数のヘッダー行を占めるが、パス要素は一つだけである。
                # 一方、横結合では親項目を配下の各列へ引き継ぐ。
                if merged.min_row < row and merged.max_row > merged.min_row:
                    return None
                return sheet.cell(merged.min_row, merged.min_col).value
        return None
    columns: list[str] = []
    for column in range(2, sheet.max_column + 1):
        parts = [str(header_value(row, column)).strip() for row in range(1, header_depth + 1) if header_value(row, column) not in (None, '')]
        # 料金.オプション.オプション のような階層間の同名項目を保持する。
        # 縦結合による重複部分は header_value() ですでに除外している。
        path_name = '.'.join(parts)
        columns.append(path_name)
    expected = set(scalar_paths(schema))
    actual = set(columns)
    unexpected = sorted(actual - expected)
    missing = sorted(expected - actual)
    if not columns or unexpected:
        details = [tr('data_excel.header_mismatch')]
        if unexpected:
            details.append(f"Excel only: {', '.join(unexpected)}")
        if missing:
            details.append(f"Current structure only: {', '.join(missing)}")
        details.append(f"Excel headers: {', '.join(columns) if columns else '(none)'}")
        raise ValueError('\n'.join(details))
    grouped: list[tuple[str, list[dict[str, Any]]]] = []
    current_name = ''
    current_rows: list[dict[str, Any]] = []
    for row_number in range(header_depth + 1, sheet.max_row + 1):
        name_value = sheet.cell(row_number, 1).value
        values = {
            path_name: strip_data_whitespace(sheet.cell(row_number, index + 2).value)
            for index, path_name in enumerate(columns)
        }
        if name_value not in (None, ''):
            if current_name:
                grouped.append((current_name, current_rows))
            current_name = str(name_value).strip()
            current_rows = []
        if current_name and any((value is not None for value in values.values())):
            current_rows.append(values)
    if current_name:
        grouped.append((current_name, current_rows))
    if not grouped:
        raise ValueError('data_excel.no_data')
    execution_settings: dict[str, tuple[str, bool, str]] = {}
    settings_names = {tr_language('data_excel.settings_sheet', language) for language in SUPPORTED_LANGUAGES}
    # 出力時と現在の UI 言語が異なっても設定シートを認識する。
    settings_name = next((name for name in workbook.sheetnames if name in settings_names), None)
    if _read_settings and settings_name is not None:
        settings_sheet = workbook[settings_name]
        has_summary_column = settings_sheet.max_column >= 4
        for row in settings_sheet.iter_rows(min_row=2, values_only=True):
            record_name = str(row[0] if row and row[0] is not None else '').strip()
            if not record_name:
                continue
            summary = str(row[1] if has_summary_column and len(row) > 1 and row[1] is not None else '').strip()
            enabled_index = 2 if has_summary_column else 1
            group_index = 3 if has_summary_column else 2
            enabled_value = row[enabled_index] if len(row) > enabled_index else True
            enabled_words = {tr_language('common.yes', language).lower() for language in SUPPORTED_LANGUAGES}
            enabled_words.add(tr('data_excel.enabled_value').lower())
            enabled = enabled_value if isinstance(enabled_value, bool) else str(enabled_value).lower() in {'true', '1', 'yes', *enabled_words}
            group = str(row[group_index] if len(row) > group_index and row[group_index] not in (None, '') else '1').strip()
            execution_settings[record_name] = (summary, enabled, group)
    owners = scalar_list_owners(schema)
    list_paths = sorted({owner for value in owners.values() for owner in value}, key=lambda value: (value.count('.'), value))
    anchors = {list_path: [column for column in columns if owners[column] and owners[column][-1] == list_path] for list_path in list_paths}
    node_by_path = {path_name: node for path_name, node in schema_paths(schema)}

    def assign(data: dict[str, Any], path_name: str, value: Any, indexes: dict[str, int]) -> None:
        # パスと行ごとの list 添字を使って、平坦なセルを階層辞書へ戻す。
        container: Any = data
        prefix = ''
        parts = path_name.split('.')
        for index, part in enumerate(parts):
            prefix = f'{prefix}.{part}' if prefix else part
            node = node_by_path[prefix]
            last = index == len(parts) - 1
            if node['type'] == 'list':
                target = container.setdefault(part, [])
                item_index = indexes[prefix]
                while len(target) <= item_index:
                    target.append(new_list_item(node))
                container = target[item_index]
            elif last:
                if node['type'] == 'number' and value not in (None, ''):
                    value = float(value) if isinstance(value, float) and (not value.is_integer()) else int(value)
                elif node['type'] == 'boolean':
                    yes_words = {tr_language('common.yes', language).lower() for language in SUPPORTED_LANGUAGES}
                    value = value if isinstance(value, bool) else str(value).lower() in {'true', '1', 'yes', *yes_words}
                container[part] = value
            elif node['type'] == 'object':
                container = container.setdefault(part, {})
    result: list[dict[str, Any]] = []
    for name, rows in grouped:
        data = {
            child['name']: record_default_value(child)
            for child in schema.get('children', [])
            if not child.get('data_template', False)
        }
        counters = {path_name: 0 for path_name in list_paths}
        previous = dict(counters)
        for row_index, values in enumerate(rows):
            previous = dict(counters)
            for list_path in list_paths:
                parents = [owner for owner in list_paths if list_path.startswith(owner + '.')]
                parent = max(parents, key=len) if parents else None
                if row_index == 0 or (parent and counters[parent] != previous[parent]):
                    counters[list_path] = 0
                elif any((values.get(column) is not None for column in anchors[list_path])):
                    counters[list_path] += 1
            for path_name, value in values.items():
                if value is not None:
                    assign(data, path_name, value, counters)
        # Excel 側で行を追加・並べ替えても、表示中のデータ名で実行設定を正しく対応させる。
        # 設定シート自体がない場合も、対象 PCL の行だけがない場合も同じ既定値を使う。
        setting = execution_settings.get(name, ('', True, '1'))
        result.append({'name': name, 'summary': setting[0], 'enabled': setting[1], 'execution_group': setting[2], 'data': strip_data_whitespace(normalize_record(schema, data))})
    return result


def _read_split_records_excel(
        path: str | Path, workbook: Any, schema: dict[str, Any], metadata: dict[str, Any],
) -> list[dict[str, Any]]:
    """各シートをデータ名で結合し、一つの PCL データへ復元する。"""
    main_schema, _sections = _split_excel_schema(schema)
    data_names = {tr_language('data_excel.data_sheet', language) for language in SUPPORTED_LANGUAGES}
    main_name = next((name for name in workbook.sheetnames if name in data_names), None)
    if main_name is None:
        raise ValueError(tr('実行データシートが見つかりません。'))
    records = read_records_excel(
        path, main_schema, workbook=workbook, _sheet=workbook[main_name],
        _allow_split=False,
    )

    def unique_by_name(rows: list[dict[str, Any]], sheet_name: str) -> dict[str, dict[str, Any]]:
        indexed: dict[str, dict[str, Any]] = {}
        for row in rows:
            name = str(row['name']).strip()
            if name in indexed:
                raise ValueError(f'「{sheet_name}{tr("」シートに重複したデータ名があります: ")}{name}')
            indexed[name] = row
        return indexed

    records_by_name = unique_by_name(records, main_name)
    nodes_by_path = {path_name: node for path_name, node in schema_paths(schema)}

    def assign_section(data: dict[str, Any], path_name: str, value: dict[str, Any]) -> None:
        container = data
        parts = path_name.split('.')
        for part in parts[:-1]:
            current = container.get(part)
            if not isinstance(current, dict):
                current = {}
                container[part] = current
            container = current
        container[parts[-1]] = value

    def remove_section(data: dict[str, Any], path_name: str) -> None:
        """未追加テンプレートを、Excel の空セルから誤って生成しない。"""
        container: Any = data
        parts = path_name.split('.')
        for part in parts[:-1]:
            if not isinstance(container, dict) or not isinstance(container.get(part), dict):
                return
            container = container[part]
        if isinstance(container, dict):
            container.pop(parts[-1], None)

    for mapping in metadata['mappings']:
        if not isinstance(mapping, dict):
            raise ValueError(tr('Excel のシート対応情報が正しくありません。'))
        sheet_name = str(mapping.get('sheet', ''))
        path_name = str(mapping.get('path', ''))
        if str(mapping.get('kind', '')) == 'template':
            template = template_by_id(schema, str(mapping.get('template_id', '')))
            if template is None or sheet_name not in workbook.sheetnames:
                continue
            simple_format = bool(mapping.get('simple'))
            template_schema = {
                'name': 'Data', 'type': 'object', 'children': [
                    *([{
                        'name': '配置先', 'type': 'text',
                    }, {
                        'name': '順番', 'type': 'number',
                    }, {
                        'name': 'テンプレート名', 'type': 'text',
                    }] if simple_format else [{
                        'name': '__instance_id__', 'type': 'text',
                    }, {
                        'name': '__instance_name__', 'type': 'text',
                    }, {
                        'name': '__container_path__', 'type': 'text',
                    }]),
                    *copy.deepcopy(template.get('children', [])),
                ],
            }
            instance_rows = read_records_excel(
                path, template_schema, workbook=workbook, _sheet=workbook[sheet_name],
                _read_settings=False, _allow_split=False,
            )
            for row in instance_rows:
                record_name = str(row.get('name', '')).strip()
                if record_name not in records_by_name:
                    raise ValueError(
                        f'「{sheet_name}」シートのデータ名がメインシートに存在しません: {record_name}'
                    )
                values = dict(row.get('data', {}))
                if simple_format:
                    destination = str(values.pop('配置先', '') or '').strip()
                    raw_position = values.pop('順番', 1)
                    instance_name = str(values.pop('テンプレート名', '') or '').strip()
                    try:
                        position = max(0, int(raw_position or 1) - 1)
                    except (TypeError, ValueError):
                        position = 0
                    instance = {
                        'instance_id': uuid.uuid4().hex,
                        'template_id': str(template.get('template_id', '')),
                        'template_name': str(template.get('name', '')).strip(),
                        'name': instance_name or str(template.get('name', '')),
                        'data': values,
                    }
                    target_data = records_by_name[record_name]['data']
                    if not destination:
                        container = target_data.setdefault(TEMPLATE_INSTANCES_KEY, [])
                    else:
                        container = _resolve_container_path(target_data, destination)
                    container.insert(min(position, len(container)), instance)
                    continue
                instance_id = str(values.pop('__instance_id__', '')).strip()
                instance_name = str(values.pop('__instance_name__', '')).strip()
                raw_path = str(values.pop('__container_path__', '')).strip()
                instance = {
                    'instance_id': instance_id,
                    'template_id': str(template.get('template_id', '')),
                    'name': instance_name or str(template.get('name', '')),
                    'data': values,
                }
                target_data = records_by_name[record_name]['data']
                try:
                    instance_path = json.loads(raw_path) if raw_path else None
                except json.JSONDecodeError:
                    instance_path = None
                if not isinstance(instance_path, list) or not instance_path:
                    # 格納パスを持たない旧形式は、従来どおり PCL 直下へ復元する。
                    target_data.setdefault(TEMPLATE_INSTANCES_KEY, []).append(instance)
                else:
                    _assign_value_at_path(target_data, instance_path, instance)
            continue
        node = nodes_by_path.get(path_name)
        # 空データの分割シートは省略できる。対応ノード自体は現在構造に必要となる。
        if node is None or node.get('type') != 'object':
            raise ValueError(f'{tr("現在のデータ構造に Excel の分割項目がありません: ")}{path_name}')
        if sheet_name not in workbook.sheetnames:
            continue
        section_schema = {
            'name': 'Data', 'type': 'object',
            'children': copy.deepcopy(node.get('children', [])),
        }
        section_rows = read_records_excel(
            path, section_schema, workbook=workbook, _sheet=workbook[sheet_name],
            _read_settings=False, _allow_split=False,
        )
        section_by_name = unique_by_name(section_rows, sheet_name)
        orphan_names = sorted(set(section_by_name) - set(records_by_name))
        if orphan_names:
            raise ValueError(
                f'「{sheet_name}{tr("」シートのデータ名が「")}{main_name}'
                f'{tr("」シートに存在しません: ")}{orphan_names[0]}'
            )
        for name, section in section_by_name.items():
            assign_section(records_by_name[name]['data'], path_name, section['data'])

    for path_name, present_names in metadata.get('template_presence', {}).items():
        node = nodes_by_path.get(path_name)
        if node is None or not node.get('data_template', False) or not isinstance(present_names, list):
            continue
        present = {str(name).strip() for name in present_names}
        for record in records:
            if record['name'] in present:
                if _value_at_path(record['data'], path_name) is None:
                    assign_section(record['data'], path_name, default_value(node))
            else:
                remove_section(record['data'], path_name)

    for record in records:
        record['data'] = strip_data_whitespace(normalize_record(schema, record['data']))
    return records

def default_value(node: dict[str, Any]) -> Any:
    kind = node['type']
    if kind in ('object', 'list'):
        value = {child['name']: default_value(child) for child in node.get('children', [])}
        return [] if kind == 'list' else value
    if kind == 'number':
        return 0
    if kind == 'boolean':
        return False
    return ''

def new_list_item(node: dict[str, Any], populate_nested_lists: bool=False) -> dict[str, Any]:

    def build(child: dict[str, Any]) -> Any:
        if child['type'] == 'list':
            return [new_list_item(child, True)] if populate_nested_lists else []
        if child['type'] == 'object':
            return {grandchild['name']: build(grandchild) for grandchild in child.get('children', [])}
        return default_value(child)
    return {child['name']: build(child) for child in node.get('children', [])}

def normalize_record(schema: dict[str, Any], data: dict[str, Any]) -> dict[str, Any]:
    result = dict(data)
    for child in schema.get('children', []):
        name = child['name']
        if name not in result:
            # 任意テンプレートは、JSON に存在するときだけ実体化する。
            if child.get('data_template', False):
                continue
            result[name] = record_default_value(child)
        elif child['type'] == 'object':
            result[name] = normalize_record(child, result[name]) if isinstance(result[name], dict) else record_default_value(child)
        elif child['type'] == 'list':
            result[name] = [
                # list に追加したテンプレート実体へ、通常項目の既定値を混在させない。
                item if (
                    isinstance(item, dict)
                    and isinstance(item.get('data'), dict)
                    and bool(str(item.get('template_id', '')).strip())
                ) else normalize_record(child, item) if isinstance(item, dict) else new_list_item(child)
                for item in result[name]
            ] if isinstance(result[name], list) else []
    return result

def record_default_value(node: dict[str, Any]) -> Any:
    """Excel 読込時も任意テンプレートを自動追加しない既定値を返す。"""
    if node['type'] == 'object':
        return {
            child['name']: record_default_value(child)
            for child in node.get('children', [])
            if not child.get('data_template', False)
        }
    return default_value(node)

def schema_paths(schema: dict[str, Any]) -> list[tuple[str, dict[str, Any]]]:
    result: list[tuple[str, dict[str, Any]]] = []

    def walk(node: dict[str, Any], prefix: str) -> None:
        for child in node.get('children', []):
            path = f"{prefix}.{child['name']}" if prefix else child['name']
            result.append((path, child))
            if child['type'] in ('object', 'list'):
                walk(child, path)
    walk(schema, '')
    return result

