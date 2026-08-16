from __future__ import annotations

import copy
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from core.excel_io import (
    normalize_record, read_records_excel, read_records_excel_with_schema, remap_data_for_schema_names, schema_name_path_map,
    strip_data_record_whitespace, strip_schema_name_whitespace, validate_schema, write_records_excel,
)
from core.data_templates import migrate_legacy_template_data, normalize_template_schema


class ExcelStyleTests(unittest.TestCase):
    def test_legacy_template_schema_and_data_are_migrated_without_wrapper_path(self) -> None:
        schema = {
            'name': 'Data', 'type': 'object', 'children': [{
                'name': 'テンプレート', 'type': 'object', 'children': [{
                    'name': '仮想商材', 'type': 'object', 'data_template': True,
                    'template_id': 'product',
                    'children': [{'name': '名称', 'type': 'text'}],
                }],
            }],
        }

        migrated_schema = normalize_template_schema(schema)
        migrated_data = migrate_legacy_template_data(
            schema, {'テンプレート': {'仮想商材': {'名称': 'A'}}},
        )

        self.assertEqual(migrated_schema['children'], [])
        self.assertEqual(migrated_schema['templates'][0]['template_id'], 'product')
        self.assertNotIn('テンプレート', migrated_data)
        self.assertEqual(migrated_data['_template_instances'][0]['data'], {'名称': 'A'})

    def test_normalize_record_does_not_materialize_optional_data_template(self) -> None:
        schema = {
            'name': 'Data', 'type': 'object', 'children': [
                {'name': 'common', 'type': 'text'},
                {'name': 'detail', 'type': 'object', 'data_template': True, 'children': [
                    {'name': 'value', 'type': 'text'},
                ]},
            ],
        }
        self.assertEqual(normalize_record(schema, {}), {'common': ''})
        self.assertEqual(
            normalize_record(schema, {'detail': {}}),
            {'common': '', 'detail': {'value': ''}},
        )

    def test_excel_round_trip_preserves_template_presence_per_record(self) -> None:
        schema = {
            'name': 'Data', 'type': 'object', 'children': [
                {'name': 'common', 'type': 'text'},
            ],
            'templates': [{
                'name': 'detail', 'type': 'object', 'template_id': 'detail-template',
                'children': [{'name': 'value', 'type': 'text'}],
            }],
        }
        records = [
            {'name': 'PCL_001', 'summary': '', 'enabled': True, 'execution_group': '1',
             'data': {'common': 'A', '_template_instances': [{
                 'instance_id': 'instance-1', 'template_id': 'detail-template',
                 'name': 'detail 1', 'data': {'value': 'included'},
             }, {
                 'instance_id': 'instance-2', 'template_id': 'detail-template',
                 'name': 'detail 2', 'data': {'value': 'included twice'},
             }]}},
            {'name': 'PCL_002', 'summary': '', 'enabled': True, 'execution_group': '1',
             'data': {'common': 'B'}},
        ]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'templates.xlsx'
            write_records_excel(path, schema, records)
            workbook = load_workbook(path, data_only=True)
            detail_names = [
                workbook['detail'].cell(row, 1).value
                for row in range(1, workbook['detail'].max_row + 1)
            ]
            self.assertNotIn('__WebFlowManager__', workbook.sheetnames)
            headers = {
                workbook['detail'].cell(row, column).value
                for row in range(1, 3) for column in range(1, 8)
            }
            self.assertTrue({'配置先', '順番', 'テンプレート名'} <= headers)
            self.assertFalse(any(str(value or '').startswith('__') for value in headers))
            # PCL 直下のテンプレートは配置先を空欄にする。
            self.assertIsNone(workbook['detail'].cell(3, 2).value)
            restored_records = read_records_excel(path, schema)

        self.assertEqual(detail_names.count('PCL_001'), 2)
        self.assertNotIn('PCL_002', detail_names)
        expected = copy.deepcopy(records)
        for index, instance in enumerate(expected[0]['data']['_template_instances']):
            instance['template_name'] = 'detail'
            instance['instance_id'] = restored_records[0]['data']['_template_instances'][index]['instance_id']
        self.assertEqual(restored_records, expected)

    def test_excel_round_trip_preserves_template_instance_inside_list(self) -> None:
        schema = {
            'name': 'Data', 'type': 'object',
            'children': [{'name': 'items', 'type': 'list', 'children': [
                {'name': 'label', 'type': 'text'},
            ]}],
            'templates': [{
                'name': 'detail', 'type': 'object', 'template_id': 'detail-template',
                'children': [{'name': 'value', 'type': 'text'}],
            }],
        }
        records = [{
            'name': 'PCL_001', 'summary': '', 'enabled': True, 'execution_group': '1',
            'data': {'items': [{
                'instance_id': 'instance-1', 'template_id': 'detail-template',
                'name': 'detail 1', 'data': {'value': 'nested'},
            }]},
        }]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'nested-template.xlsx'
            write_records_excel(path, schema, records)
            workbook = load_workbook(path, data_only=True)
            self.assertNotIn('__WebFlowManager__', workbook.sheetnames)
            restored_records = read_records_excel(path, schema)

        expected = copy.deepcopy(records)
        expected[0]['data']['items'][0]['template_name'] = 'detail'
        expected[0]['data']['items'][0]['instance_id'] = restored_records[0]['data']['items'][0]['instance_id']
        self.assertEqual(restored_records, expected)

    def test_excel_sheet_setting_rejects_scalar_and_nested_split_nodes(self) -> None:
        with self.assertRaises(ValueError):
            validate_schema({
                'name': 'Data', 'type': 'object', 'children': [
                    {'name': '値', 'type': 'text', 'excel_sheet': True},
                ],
            })
        with self.assertRaises(ValueError):
            validate_schema({
                'name': 'Data', 'type': 'object', 'children': [{
                    'name': '親', 'type': 'object', 'excel_sheet': True, 'children': [{
                        'name': '子', 'type': 'object', 'excel_sheet': True, 'children': [],
                    }],
                }],
            })

    def test_split_excel_sheets_round_trip_with_visible_record_names(self) -> None:
        schema = {
            'name': 'Data', 'type': 'object', 'children': [
                {'name': 'PCL_NO', 'type': 'text'},
                {'name': '商材別', 'type': 'object', 'children': [
                    {'name': '仮想商材1', 'type': 'object', 'excel_sheet': True, 'children': [
                        {'name': 'プラン名', 'type': 'text'},
                        {'name': 'オプション', 'type': 'list', 'children': [
                            {'name': '名称', 'type': 'text'},
                        ]},
                    ]},
                ]},
            ],
        }
        records = [{
            'name': 'PCL_001', 'summary': '概要', 'enabled': True,
            'execution_group': '2', 'data': {
                'PCL_NO': 'PCL_001',
                '商材別': {'仮想商材1': {
                    'プラン名': '標準',
                    'オプション': [{'名称': 'A'}, {'名称': 'B'}],
                }},
            },
        }]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'split.xlsx'
            write_records_excel(path, schema, records)
            workbook = load_workbook(path, data_only=True)
            main_sheet = workbook.worksheets[0]
            self.assertIn('仮想商材1', workbook.sheetnames)
            self.assertNotIn('__WebFlowManager__', workbook.sheetnames)
            # 分割先では親パスを繰り返さず、短いヘッダーだけを表示する。
            split_headers = {
                workbook['仮想商材1'].cell(row, column).value
                for row in range(1, 3) for column in range(2, 4)
            }
            self.assertNotIn('商材別', split_headers)
            restored_records = read_records_excel(path, schema)

        self.assertEqual(restored_records, records)

    def test_split_excel_accepts_new_records_added_with_visible_names(self) -> None:
        schema = {
            'name': 'Data', 'type': 'object', 'children': [
                {'name': '共通', 'type': 'text'},
                {'name': '詳細', 'type': 'object', 'excel_sheet': True, 'children': [
                    {'name': '値', 'type': 'text'},
                ]},
            ],
        }
        records = [{
            'name': 'PCL_001', 'enabled': True, 'execution_group': '1',
            'data': {'共通': 'A', '詳細': {'値': '1'}},
        }]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'bulk-add.xlsx'
            write_records_excel(path, schema, records)
            workbook = load_workbook(path)
            workbook.worksheets[0].append(['PCL_002', 'B'])
            workbook['詳細'].append(['PCL_002', '2'])
            workbook.save(path)
            restored = read_records_excel(path, schema)

        self.assertEqual([record['name'] for record in restored], ['PCL_001', 'PCL_002'])
        self.assertEqual(restored[1]['data'], {'共通': 'B', '詳細': {'値': '2'}})

    def test_split_excel_rejects_orphan_record_names(self) -> None:
        schema = {
            'name': 'Data', 'type': 'object', 'children': [
                {'name': '共通', 'type': 'text'},
                {'name': '詳細', 'type': 'object', 'excel_sheet': True, 'children': [
                    {'name': '値', 'type': 'text'},
                ]},
            ],
        }
        records = [{
            'name': 'PCL_001', 'enabled': True, 'execution_group': '1',
            'data': {'共通': 'A', '詳細': {'値': '1'}},
        }]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'orphan.xlsx'
            write_records_excel(path, schema, records)
            workbook = load_workbook(path)
            workbook['詳細'].append(['PCL_999', 'orphan'])
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, 'PCL_999'):
                read_records_excel(path, schema)

    def test_empty_split_sheet_is_omitted_by_default_and_can_be_forced(self) -> None:
        schema = {
            'name': 'Data', 'type': 'object', 'children': [
                {'name': '共通', 'type': 'text'},
                {'name': '詳細', 'type': 'object', 'excel_sheet': True, 'children': [
                    {'name': '値', 'type': 'text'},
                ]},
            ],
        }
        records = [{
            'name': 'PCL_001', 'summary': '', 'enabled': True, 'execution_group': '1',
            'data': {'共通': 'A', '詳細': {'値': ''}},
        }]
        with tempfile.TemporaryDirectory() as directory:
            omitted_path = Path(directory) / 'omitted.xlsx'
            write_records_excel(omitted_path, schema, records)
            workbook = load_workbook(omitted_path, data_only=True)
            self.assertNotIn('詳細', workbook.sheetnames)
            self.assertNotIn('__WebFlowManager__', workbook.sheetnames)
            restored_records = read_records_excel(omitted_path, schema)

            forced_schema = {
                **schema,
                'children': [schema['children'][0], {
                    **schema['children'][1], 'excel_skip_empty': False,
                }],
            }
            forced_path = Path(directory) / 'forced.xlsx'
            write_records_excel(forced_path, forced_schema, records)
            forced_workbook = load_workbook(forced_path, data_only=True)

        self.assertEqual(restored_records, records)
        self.assertIn('詳細', forced_workbook.sheetnames)

    def test_excel_headers_can_restore_schema_and_records(self) -> None:
        schema = {
            'name': 'Data', 'type': 'object', 'children': [
                {'name': 'profile', 'type': 'object', 'children': [
                    {'name': 'name', 'type': 'text'},
                    {'name': 'active', 'type': 'boolean'},
                    {'name': 'count', 'type': 'number'},
                ]},
                {'name': 'items', 'type': 'list', 'children': [
                    {'name': 'code', 'type': 'text'},
                ]},
            ],
        }
        records = [{
            'name': 'PCL_001', 'summary': 'summary', 'enabled': True,
            'execution_group': '2', 'data': {
                'profile': {'name': 'Alice', 'active': True, 'count': 3},
                'items': [{'code': 'A'}, {'code': 'B'}],
            },
        }]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'schema-import.xlsx'
            write_records_excel(path, schema, records)
            restored_schema, restored_records = read_records_excel_with_schema(path)

        self.assertEqual(restored_schema, schema)
        self.assertEqual(restored_records, records)

    def test_schema_names_and_existing_data_keys_are_trimmed_together(self) -> None:
        schema = {
            'name': '\u3000Data\n', 'type': 'object',
            'children': [{
                'name': '\t料金 \u3000', 'type': 'object',
                'children': [{'name': '\nオプション\r', 'type': 'text'}],
            }],
        }
        cleaned = strip_schema_name_whitespace(schema)
        migrated = remap_data_for_schema_names(
            schema, cleaned, {'\t料金 \u3000': {'\nオプション\r': '値'}},
        )

        self.assertEqual(cleaned['name'], 'Data')
        self.assertEqual(cleaned['children'][0]['name'], '料金')
        self.assertEqual(cleaned['children'][0]['children'][0]['name'], 'オプション')
        self.assertEqual(migrated, {'料金': {'オプション': '値'}})
        self.assertEqual(schema_name_path_map(schema, cleaned), {
            '\t料金 \u3000': '料金',
            '\t料金 \u3000.\nオプション\r': '料金.オプション',
        })

    def test_schema_name_trim_rejects_resulting_sibling_duplicate(self) -> None:
        schema = {
            'name': 'Data', 'type': 'object',
            'children': [
                {'name': '料金', 'type': 'text'},
                {'name': '料金\u3000', 'type': 'text'},
            ],
        }
        with self.assertRaises(ValueError):
            strip_schema_name_whitespace(schema)

    def test_data_json_record_strips_unicode_whitespace(self) -> None:
        cleaned = strip_data_record_whitespace({
            'name': '\u3000PCL_001\n',
            'summary': '\t概要\r',
            'execution_group': '\n 1 \u3000',
            'enabled': True,
            'data': {'outer': {'text': '\r\n 値 \t'}, 'items': ['\u3000A ', '\nB\r']},
        })

        self.assertEqual(cleaned['name'], 'PCL_001')
        self.assertEqual(cleaned['summary'], '概要')
        self.assertEqual(cleaned['execution_group'], '1')
        self.assertEqual(cleaned['data'], {'outer': {'text': '値'}, 'items': ['A', 'B']})

    def test_data_management_excel_strips_unicode_whitespace(self) -> None:
        schema = {
            'name': 'Data', 'type': 'object',
            'children': [{
                'name': 'items', 'type': 'list',
                'children': [{'name': 'value', 'type': 'text'}],
            }],
        }
        records = [{
            'name': '\u3000PCL_001\n',
            'summary': '\t概要\r\n',
            'enabled': True,
            'execution_group': '\u30001\t',
            'data': {'items': [{'value': '\n 値 \u3000'}, {'value': '\t次\r'}]},
        }]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'trimmed.xlsx'
            write_records_excel(path, schema, records)
            workbook = load_workbook(path, data_only=True)
            data_sheet = workbook.worksheets[0]
            settings_sheet = workbook.worksheets[1]
            self.assertEqual(data_sheet.cell(3, 1).value, 'PCL_001')
            self.assertEqual(data_sheet.cell(3, 2).value, '値')
            self.assertEqual(data_sheet.cell(4, 2).value, '次')
            self.assertEqual(settings_sheet.cell(2, 2).value, '概要')
            self.assertEqual(settings_sheet.cell(2, 4).value, '1')
            restored = read_records_excel(path, schema)

        self.assertEqual(restored[0]['name'], 'PCL_001')
        self.assertEqual(restored[0]['summary'], '概要')
        self.assertEqual(restored[0]['data']['items'], [{'value': '値'}, {'value': '次'}])

    def test_import_uses_named_data_sheet_when_settings_sheet_is_active(self) -> None:
        """実行設定を表示したまま保存しても、データ一覧を正しく読み込む。"""
        schema = {
            'name': 'Data', 'type': 'object',
            'children': [{'name': 'value', 'type': 'text'}],
        }
        records = [{
            'name': 'PCL_001', 'summary': '概要', 'enabled': True,
            'execution_group': '1', 'data': {'value': '値'},
        }]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'active-settings.xlsx'
            write_records_excel(path, schema, records)
            workbook = load_workbook(path)
            workbook.active = 1
            workbook.save(path)

            restored = read_records_excel(path, schema)

        self.assertEqual(restored, records)

    def test_execution_settings_are_optional_and_missing_records_use_defaults(self) -> None:
        schema = {
            'name': 'Data', 'type': 'object',
            'children': [{'name': 'value', 'type': 'text'}],
        }
        records = [{
            'name': 'PCL_001', 'summary': '設定あり', 'enabled': False,
            'execution_group': '3', 'data': {'value': 'A'},
        }, {
            'name': 'PCL_002', 'summary': '削除対象', 'enabled': False,
            'execution_group': '4', 'data': {'value': 'B'},
        }]
        with tempfile.TemporaryDirectory() as directory:
            partial_path = Path(directory) / 'partial-settings.xlsx'
            write_records_excel(partial_path, schema, records)
            workbook = load_workbook(partial_path)
            settings = workbook['実行設定']
            settings.delete_rows(3)
            workbook.save(partial_path)
            partial = read_records_excel(partial_path, schema)

            no_settings_path = Path(directory) / 'no-settings.xlsx'
            workbook = load_workbook(partial_path)
            del workbook['実行設定']
            workbook.save(no_settings_path)
            without_settings = read_records_excel(no_settings_path, schema)

        self.assertEqual(
            (partial[0]['summary'], partial[0]['enabled'], partial[0]['execution_group']),
            ('設定あり', False, '3'),
        )
        self.assertEqual(
            (partial[1]['summary'], partial[1]['enabled'], partial[1]['execution_group']),
            ('', True, '1'),
        )
        self.assertTrue(all(
            (record['summary'], record['enabled'], record['execution_group']) == ('', True, '1')
            for record in without_settings
        ))

    def test_repeated_names_in_multilevel_header_round_trip(self) -> None:
        schema = {
            'name': 'Data',
            'type': 'object',
            'children': [{
                'name': '料金',
                'type': 'object',
                'children': [{
                    'name': 'オプション',
                    'type': 'list',
                    'children': [{'name': 'オプション', 'type': 'text'}],
                }],
            }],
        }
        records = [{
            'name': 'PCL_001',
            'enabled': True,
            'execution_group': '1',
            'data': {'料金': {'オプション': [
                {'オプション': 'A'}, {'オプション': 'B'},
            ]}},
        }]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'repeated-header.xlsx'
            write_records_excel(path, schema, records)
            restored = read_records_excel(path, schema)

        self.assertEqual(restored[0]['data'], records[0]['data'])

    def test_header_mismatch_reports_actual_differences(self) -> None:
        schema = {
            'name': 'Data', 'type': 'object',
            'children': [{'name': '料金', 'type': 'text'}],
        }
        records = [{
            'name': 'PCL_001', 'enabled': True, 'execution_group': '1',
            'data': {'料金': '100'},
        }]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'bad-header.xlsx'
            write_records_excel(path, schema, records)
            workbook = load_workbook(path)
            workbook.worksheets[0].cell(1, 2, '金額')
            workbook.save(path)
            with self.assertRaises(ValueError) as raised:
                read_records_excel(path, schema)

        message = str(raised.exception)
        self.assertIn('Excel only: 金額', message)
        self.assertIn('Current structure only: 料金', message)
        self.assertIn('Excel headers: 金額', message)

    def test_all_values_are_left_aligned_and_header_levels_have_distinct_colors(self) -> None:
        schema = {
            'name': 'Data',
            'type': 'list',
            'children': [
                {'name': 'PCL_NO', 'type': 'text'},
                {
                    'name': 'output',
                    'type': 'object',
                    'children': [
                        {'name': 'contract_id', 'type': 'text'},
                        {'name': 'quantity', 'type': 'number'},
                    ],
                },
            ],
        }
        records = [{
            'name': 'PCL_001',
            'enabled': True,
            'execution_group': '1',
            'data': {'PCL_NO': 'PCL_001', 'output': {'contract_id': 'C-001', 'quantity': 12}},
        }]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'records.xlsx'
            write_records_excel(path, schema, records)
            workbook = load_workbook(path)

        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        self.assertEqual(cell.alignment.horizontal, 'left', f'{sheet.title}!{cell.coordinate}')

        data_sheet = workbook.worksheets[0]
        top_color = data_sheet.cell(1, 3).fill.fgColor.rgb
        child_color = data_sheet.cell(2, 3).fill.fgColor.rgb
        self.assertNotEqual(top_color, child_color)
        self.assertEqual(top_color, '002F5597')
        self.assertEqual(child_color, '004472C4')


if __name__ == '__main__':
    unittest.main()
